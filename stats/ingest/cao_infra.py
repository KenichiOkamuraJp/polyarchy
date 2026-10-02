"""
内閣府「社会資本ストック推計」データ（Excel）からの取込（第 12 弾 第 6 便・2026-10-02）。

- 要件源＝利用側の財政 PL・BS 分析（下水道など地方の一般会計等の外にある社会資本を見積もる）。
- 取得元＝https://www5.cao.go.jp/keizai2/ioj/result/ioj_data.html の年度版（data/stock_fy.xlsx・pref_{gross,prod,net,flow}_fy.xlsx）。
  内閣府サイトの利用規約＝政府標準利用規約準拠（◎・出典明示）。暦年版（*_cy）とストック毀損額は収録しない。
- 全国（layout=national）＝部門ごとに 1 シート（17部門計＋部門・部門の内訳 21）。行＝年度（A 列の西暦）・列＝2 段見出し
  （2 行目＝名目投資額／デフレーター／実質投資額／ストック（2015暦年基準）＝右へ引き継ぐ・3 行目＝新設改良費／災害復旧費／粗／純／生産的）。
  列は (2 行目, 3 行目) の組がちょうど 1 列に当たることを確かめて決める。
- 都道府県（layout=pref）＝部門ごとに 1 シート（16部門計＋19＝鉄道を除く）。行＝都道府県（「東京」等の短い名前→JIS コード）と「全国」・列＝年度（3 行目）。
  47 都道府県＋全国がそろわなければエラー。'－'（沖縄の復帰前 等）は値ではない。
- 金額は 2015 年価格（実質）・百万円。名目投資額だけが名目。値は表示どおりの整数（デフレーターは小数 3 桁）。
- accessor: {"type": "cao_infra_xlsx", "file": "stock_fy.xlsx", "layout": "national", "sheet": "17部門計", "group": "ストック（2015暦年基準）", "sub": "純"}
            {"type": "cao_infra_xlsx", "file": "pref_net_fy.xlsx", "layout": "pref", "sheet": "16部門計"}
- 原本は stats/data/cache/cao/<取得日>/infra/。

実行（リポジトリ root）：  python -m stats.ingest.cao_infra --all
"""
from __future__ import annotations

import sys
from typing import Optional

from polyarchy_common.logsetup import configure_quiet_logging, get_logger

from stats.core.dim_vocab import PREF_NAMES, PREF_CODE
from stats.core.registry import Series
from stats.core.values import ValueRecord
from stats.ingest._base import SourceError, _norm, cache_rel, cell_text, fetch, finish, run_cli, today

configure_quiet_logging()
log = get_logger("polyarchy.stats.ingest.cao_infra")
BASE = "https://www5.cao.go.jp/keizai2/ioj/result/data/"
# 原表の短い名前（「東京」「大阪」「北海道」）→ JIS コード。都道府県名から「都・府・県」を落とした形（北海道はそのまま）
SHORT_PREF = {(n if n == "北海道" else n[:-1]): PREF_CODE[n] for n in PREF_NAMES}


class CaoInfraError(SourceError):
    pass


_WB: dict[tuple[str, str], tuple[object, str, str]] = {}


def _workbook(file: str, day: str):
    import openpyxl  # 取込時のみ
    key = (file, day)
    if key not in _WB:
        cp, published = fetch(BASE + file, day=day, kind="cao", subdir="infra", skip_if_exists=True)
        _WB[key] = (openpyxl.load_workbook(cp, data_only=True), cache_rel(cp), published)
    return _WB[key]


def _text(c) -> Optional[str]:
    v = c.value
    if isinstance(v, bool) or v is None or isinstance(v, str):
        return None  # '－' 等は値ではない
    if isinstance(v, int):
        return str(v)
    return cell_text(v, c.number_format)


def _national(s: Series, ws, cache: str, published: str, day: str) -> list[ValueRecord]:
    a = s.accessor
    group, sub = _norm(a["group"]), _norm(a["sub"])
    cols, carried = [], ""
    for c in range(2, ws.max_column + 1):
        g = _norm(ws.cell(2, c).value)
        carried = g or carried
        if carried == group and _norm(ws.cell(3, c).value) == sub:
            cols.append(c)
    if len(cols) != 1:
        raise CaoInfraError(f"{s.series_id}: {ws.title} で見出し ({a['group']}, {a['sub']}) が {len(cols)} 列")
    col = cols[0]
    recs = []
    for r in range(5, ws.max_row + 1):
        y = ws.cell(r, 1).value
        if not isinstance(y, int):
            continue
        text = _text(ws.cell(r, col))
        if text is None:
            continue
        recs.append(ValueRecord(series_id=s.series_id, period=f"FY{y}", region="JP", value=text, status="", vintage=published or day,
                                retrieved_at=day, published_at=published,
                                accessor={"type": "cao_infra_xlsx", "file": a["file"], "sheet": ws.title, "group": a["group"], "sub": a["sub"],
                                          "cell": ws.cell(r, col).coordinate, "cache": cache}))
    return recs


def _pref(s: Series, ws, cache: str, published: str, day: str) -> list[ValueRecord]:
    a = s.accessor
    years = {c: ws.cell(3, c).value for c in range(2, ws.max_column + 1) if isinstance(ws.cell(3, c).value, int)}
    if _norm(ws.cell(3, 1).value) != "年度" or not years:
        raise CaoInfraError(f"{s.series_id}: {ws.title} の 3 行目が年度の見出しではない")
    rows: dict[str, int] = {}
    for r in range(4, ws.max_row + 1):
        name = _norm(ws.cell(r, 1).value)
        code = "JP" if name == "全国" else SHORT_PREF.get(name)
        if code:
            if code in rows:
                raise CaoInfraError(f"{s.series_id}: {ws.title} で {name} が 2 行")
            rows[code] = r
    if len(rows) != 48:
        raise CaoInfraError(f"{s.series_id}: {ws.title} の都道府県＋全国が {len(rows)} 行（48 ではない）")
    recs = []
    for code, r in rows.items():
        for c, y in years.items():
            text = _text(ws.cell(r, c))
            if text is None:
                continue
            recs.append(ValueRecord(series_id=s.series_id, period=f"FY{y}", region=code, value=text, status="", vintage=published or day,
                                    retrieved_at=day, published_at=published,
                                    accessor={"type": "cao_infra_xlsx", "file": a["file"], "sheet": ws.title,
                                              "cell": ws.cell(r, c).coordinate, "cache": cache}))
    return recs


def ingest_series(s: Series, *, dry_run: bool = False, day: Optional[str] = None) -> int:
    day = day or today()
    a = s.accessor
    if a.get("type") != "cao_infra_xlsx":
        raise ValueError(f"{s.series_id}: accessor.type={a.get('type')} は cao_infra 取込の対象外")
    wb, cache, published = _workbook(a["file"], day)
    if a["sheet"] not in wb.sheetnames:
        raise CaoInfraError(f"{s.series_id}: シート {a['sheet']!r} が無い（{wb.sheetnames}）")
    ws = wb[a["sheet"]]
    recs = (_national if a["layout"] == "national" else _pref)(s, ws, cache, published, day)
    log.info("%s: %s!%s → 値 %d", s.series_id, a["file"], a["sheet"], len(recs))
    return finish(s, recs, dry_run=dry_run, exc=CaoInfraError)


def main(argv: Optional[list[str]] = None) -> int:
    return run_cli(("cao_infra_xlsx",), ingest_series, "内閣府 社会資本ストック推計 Excel 取込", argv)


if __name__ == "__main__":
    sys.exit(main())
