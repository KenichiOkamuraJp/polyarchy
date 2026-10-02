"""
財務省「国の財務書類」（一般会計・特別会計の合算＝national／連結＝consolidated）Excel 版からの取込（第 12 弾 第 3 便・2026-10-02）。

- 要件源＝利用側の財政 PL・BS 分析（国の財務書類と SNA 一般政府の差の分解。利用側は PDF を手で読んでいた＝単体 PDF は数字が特殊グリフで化ける）。
- 取得元＝https://www.mof.go.jp/policy/budget/report/public_finance_fact_sheet/fy{Y}/national/fy{Y}{gassan|renketsu}.xlsx
  （年版 Y＝会計年度。2026-10-02 に FY2018〜FY2024 の頁から実在を確認＝それより前は国立国会図書館のアーカイブのみ＝取りに行かない）。
  財務省サイトの利用規約＝公共データ利用規約（第1.0版・PDL1.0）（◎・出典明示）＝登録簿の license と同じ。
- 表＝各ファイルの先頭 4 シート（貸借対照表・業務費用計算書・資産・負債差額増減計算書・区分別収支計算書。連結は「連結」が前に付く）。
  列＝[ラベル, 前会計年度, 本会計年度]（貸借対照表は左右 2 面＝資産の部 A–C・負債の部 D–F）。単位は百万円（整数）。
  ラベルは正規化（空白除去）し、**先頭のローマ数字（Ⅰ〜Ⅸ）を外して照合**する＝連結 FY2021・FY2022 は「連結範囲の変動に伴う増減」が挟まり番号が 1 つずれる。
  1 シートで A 列・D 列を通して**ちょうど 1 行**に一致しなければエラー（黙って選ばない）。その年版にラベルが無ければその年は値なし（ログに残す）。
- **値の採り方**＝その年度を本年度とする年版の「本会計年度」列（原本の公表）。その年度の年版が無い（または読めない）ときだけ、翌年版の「前会計年度」列で埋める
  （accessor.col に前会計年度と刻む）。見出し（3 行目の「前会計年度」「本会計年度」・4 行目の和暦の年）が年版と合わなければエラー。
  ★ FY2024 の一般会計・特別会計（と一般会計）の Excel 版は**暗号化されたファイル**（OLE の EncryptedPackage・2026-10-02 確認）＝解かない＝読めない年版として扱う
  （FY2025 版が出れば前会計年度列で FY2024 が埋まる）。
- accessor: {"type": "mof_zaimu_xlsx", "kind": "national"|"consolidated", "sheet": 0..3, "row_label": "公債", "editions": [2018, …]}
  （keep_numeral＝ローマ数字を外さずに照合＝区分別収支計算書の見出し「Ⅰ 業務収支」と合計行「業務収支」を取り違えない）
- 原本は stats/data/cache/mof/<取得日>/zaimu/。

実行（リポジトリ root）：  python -m stats.ingest.mof_zaimu --all
"""
from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import Optional

from polyarchy_common.logsetup import configure_quiet_logging, get_logger

from stats.core.periods import ERA_BASE
from stats.core.registry import Series
from stats.core.values import ValueRecord
from stats.ingest._base import SourceError, _norm, cache_rel, cell_text, fetch, finish, run_cli, today

configure_quiet_logging()
log = get_logger("polyarchy.stats.ingest.mof_zaimu")
BASE = "https://www.mof.go.jp/policy/budget/report/public_finance_fact_sheet/"
FILE = {"national": "gassan", "consolidated": "renketsu"}
SHEET_WORD = ("貸借対照表", "業務費用計算書", "資産・負債差額増減計算書", "区分別収支計算書")
_ROMAN = re.compile(r"^[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]+")
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class MofZaimuError(SourceError):
    pass


def url_of(kind: str, year: int) -> str:
    return f"{BASE}fy{year}/national/fy{year}{FILE[kind]}.xlsx"


def _label(v: object) -> str:
    return _ROMAN.sub("", _norm(v))


_WB: dict[tuple[str, int, str], tuple[object, Path, str]] = {}  # 1 回の実行内のメモ（(kind, year, day) → (workbook, キャッシュ, 公表日)）


def _workbook(kind: str, year: int, day: str):
    """年版の Excel を取得して開く。暗号化されたファイル（OLE）は None（読めない年版）。"""
    import openpyxl  # 取込時のみ
    key = (kind, year, day)
    if key not in _WB:
        cp, published = fetch(url_of(kind, year), day=day, kind="mof", subdir="zaimu", skip_if_exists=True)
        if cp.read_bytes()[:8] == _OLE_MAGIC:
            log.info("%s FY%d: 暗号化されたファイル（OLE）＝読まない", kind, year)
            _WB[key] = (None, cp, published)
        else:
            _WB[key] = (openpyxl.load_workbook(cp, data_only=True), cp, published)
    return _WB[key]


def _era_year(text: str) -> Optional[int]:
    m = re.search(r"(平成|令和)(\d+|元)年", _norm(text).translate(str.maketrans("０１２３４５６７８９", "0123456789")))
    if not m:
        return None
    return ERA_BASE[m.group(1)] + (1 if m.group(2) == "元" else int(m.group(2)))


def _read(s: Series, kind: str, year: int, day: str) -> Optional[dict]:
    """年版 year のシートから {本会計年度: (値, セル), 前会計年度: (値, セル)}。暗号化・ラベル無しは None。"""
    acc = s.accessor
    wb, cp, published = _workbook(kind, year, day)
    if wb is None:
        return None
    idx = int(acc["sheet"])
    ws = wb.worksheets[idx]
    if SHEET_WORD[idx] not in ws.title:
        raise MofZaimuError(f"{s.series_id}: FY{year} の {idx} 枚目のシート名 {ws.title!r} が {SHEET_WORD[idx]} ではない")
    # keep_numeral＝ローマ数字を外さずに照合（区分別収支計算書の「Ⅰ 業務収支」は見出し・「業務収支」は合計行＝外すと 2 行に当たる）
    key = _norm if acc.get("keep_numeral") else _label
    want = key(acc["row_label"])
    hits = [(r, c) for r in range(1, ws.max_row + 1) for c in (1, 4) if key(ws.cell(r, c).value) == want]
    if not hits:
        log.info("%s: FY%d にラベル %r が無い（その年は値なし）", s.series_id, year, acc["row_label"])
        return None
    if len(hits) != 1:
        raise MofZaimuError(f"{s.series_id}: FY{year} でラベル {acc['row_label']!r} が {len(hits)} 行（1 行に確定しない）")
    row, lc = hits[0]
    out = {}
    for off, head in ((1, "前会計年度"), (2, "本会計年度")):
        col = lc + off
        if _norm(ws.cell(3, col).value) != head:
            raise MofZaimuError(f"{s.series_id}: FY{year} の 3 行目の見出しが {head} ではない（{ws.cell(3, col).value!r}）")
        # 4 行目の和暦：貸借対照表＝年度末の暦年（Y+1 年 3 月 31 日）・他の 3 表＝年度初めの暦年（自 Y 年 4 月 1 日）
        want_y = (year if head == "本会計年度" else year - 1) + (1 if idx == 0 else 0)
        if _era_year(str(ws.cell(4, col).value)) != want_y:
            raise MofZaimuError(f"{s.series_id}: FY{year} の 4 行目の年 {ws.cell(4, col).value!r} が {want_y} 年ではない")
        c = ws.cell(row, col)
        v = c.value
        text = str(v) if isinstance(v, int) and not isinstance(v, bool) else cell_text(v, c.number_format)
        if text is not None:
            out[head] = (text, c.coordinate)
    out["_cache"], out["_published"] = cache_rel(cp), published
    return out


def ingest_series(s: Series, *, dry_run: bool = False, day: Optional[str] = None) -> int:
    day = day or today()
    acc = s.accessor
    if acc.get("type") != "mof_zaimu_xlsx":
        raise ValueError(f"{s.series_id}: accessor.type={acc.get('type')} は mof_zaimu 取込の対象外")
    kind = acc["kind"]
    editions = sorted(int(y) for y in acc["editions"])
    got = {y: _read(s, kind, y, day) for y in editions}
    recs: list[ValueRecord] = []
    for fy in range(editions[0] - 1, editions[-1] + 1):
        own, nxt = got.get(fy), got.get(fy + 1)
        if own and "本会計年度" in own:
            src, head = own, "本会計年度"
            ed = fy
        elif nxt and "前会計年度" in nxt:
            src, head = nxt, "前会計年度"  # 年版が無い・読めない年度＝翌年版の前年度列（公表値）
            ed = fy + 1
        else:
            continue
        text, cell = src[head]
        recs.append(ValueRecord(
            series_id=s.series_id, period=f"FY{fy}", region="JP", value=text, status="",
            vintage=src["_published"] or day, retrieved_at=day, published_at=src["_published"],
            accessor={"type": "mof_zaimu_xlsx", "kind": kind, "file": url_of(kind, ed).removeprefix(BASE), "sheet": int(acc["sheet"]),
                      "row_label": acc["row_label"], "col": head, "cell": cell, "edition": f"FY{ed}", "cache": src["_cache"]}))
    log.info("%s: 年版 %s → 値 %d", s.series_id, editions, len(recs))
    if not recs:
        raise MofZaimuError(f"{s.series_id}: 値が 0 件（ラベル {acc['row_label']!r} がどの年版にも無い）")
    return finish(s, recs, dry_run=dry_run, exc=MofZaimuError)


def main(argv: Optional[list[str]] = None) -> int:
    return run_cli(("mof_zaimu_xlsx",), ingest_series, "財務省 国の財務書類 Excel 取込（原本キャッシュ→値ストア）", argv)


if __name__ == "__main__":
    sys.exit(main())
