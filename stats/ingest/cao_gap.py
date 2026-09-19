"""
内閣府「ＧＤＰギャップ、潜在成長率」（月例経済報告 その他の資料）の取込（第 9 弾 段 4・2026-08-29）。

- 原典＝https://www5.cao.go.jp/keizai3/getsurei/{QE公表回}gap.xlsx（例 2612gap.xlsx）。
  **ファイル名が QE の公表回に連動して毎回変わる**（更新は 2 次速報後の年 4 回）＝URL を registry に固定しない。
  取込のたびに索引ページ（getsurei-index.html）から `(\\d{4})gap.xlsx` を**ちょうど 1 件**発見する（0 件・複数件はエラー＝黙って選ばない）。
  発見した実 URL・公表回は値の accessor に刻む（原典の同じセルに戻れる）。
- シート＝四半期（1994Q1〜・年キャリー・四半期はローマ数字Ⅰ〜Ⅳ）／暦年・年度（1994〜）。
  列は見出し行（5 行目）の完全一致で 1 列に確定。値はセルの表示書式（0.0）の桁で文字列化（cell_text）。
- 潜在成長率と寄与度の 1994 期は原典が「-」＝値なし（欠測として取り込まない）。
- **※ 印**（ＧＤＰギャップ列の右の補助列）＝「0.0 は −0.0 の意味」＝値は表示どおり "0.0"・status に「※（-0.0）」を刻む。
- 日銀の需給ギャップ・潜在成長率（boj.gap）とは**推計主体・手法・頻度が異なる別系列**（混ぜない）。

実行（リポジトリ root）：
    python -m stats.ingest.cao_gap --all
    python -m stats.ingest.cao_gap --series cao.cao_gap.gdp_gap.q
"""
from __future__ import annotations

import re
import sys
from typing import Optional

from polyarchy_common.logsetup import configure_quiet_logging, get_logger

from stats.core.registry import Series
from stats.core.values import ValueRecord
from stats.ingest._base import SourceError, _norm, cache_rel, cell_text, fetch, finish, run_cli, today

configure_quiet_logging()
log = get_logger("polyarchy.stats.ingest.cao_gap")

BASE = "https://www5.cao.go.jp/keizai3/getsurei/"
INDEX_URL = BASE + "getsurei-index.html"
TYPES = ("cao_gap_xlsx",)
_ROMAN = {"Ⅰ": "Q1", "Ⅱ": "Q2", "Ⅲ": "Q3", "Ⅳ": "Q4"}


class CaoGapSourceError(SourceError):
    pass


def find_gap_file(html: str) -> tuple[str, str]:
    """索引 HTML から (ファイル名, 公表回) をちょうど 1 件確定する（純関数）。"""
    names = sorted(set(re.findall(r'href="((\d{4})gap\.xlsx)"', html)))
    if len(names) != 1:
        raise CaoGapSourceError(f"gap.xlsx のリンクが {len(names)} 件（1 件に確定しない）: {[n for n, _ in names]}")
    return names[0][0], names[0][1]


def discover_file(day: str) -> tuple[str, str]:
    """索引ページを取得して最新の gap.xlsx ファイル名と公表回を返す。"""
    cp, _ = fetch(INDEX_URL, day=day, kind="cao_gap")
    return find_gap_file(cp.read_text(encoding="utf-8", errors="replace"))


def period_of(sheet: str, year: int, qlabel: str) -> Optional[str]:
    """シート種別と行の (年, 四半期ラベル) → 正規期間表記（純関数）。四半期ラベル表記外は None。"""
    if sheet == "四半期":
        q = _ROMAN.get((qlabel or "").strip())
        return f"{year}{q}" if q else None
    return f"FY{year}" if sheet == "年度" else str(year)


def ingest_series(s: Series, *, dry_run: bool = False, day: Optional[str] = None) -> int:
    import openpyxl

    day = day or today()
    acc = s.accessor
    if acc.get("type") != "cao_gap_xlsx":
        raise ValueError(f"{s.series_id}: accessor.type={acc.get('type')} は対象外")
    fname, edition = discover_file(day)
    cp, published = fetch(BASE + fname, day=day, kind="cao_gap")
    wb = openpyxl.load_workbook(cp, data_only=True)
    sheet = str(acc.get("sheet", ""))
    if sheet not in wb.sheetnames:
        raise CaoGapSourceError(f"{s.series_id}: シート {sheet!r} が無い（{wb.sheetnames}）")
    ws = wb[sheet]
    hrow, first_row = 5, 7
    want = _norm(str(acc.get("col_header", "")))
    cols = [c for c in range(1, ws.max_column + 1) if _norm(str(ws.cell(hrow, c).value or "")) == want]
    if len(cols) != 1:
        raise CaoGapSourceError(f"{s.series_id}: 列見出し {acc.get('col_header')!r} が {len(cols)} 列（1 列に確定しない）")
    col = cols[0]
    flag_col = col + 1 if want == _norm("ＧＤＰギャップ") else 0  # ※ 印（-0.0 の表記）は ギャップ列の右
    quarterly = sheet == "四半期"
    recs: list[ValueRecord] = []
    year: Optional[int] = None
    for r in range(first_row, ws.max_row + 1):
        y = ws.cell(r, 1).value
        if isinstance(y, int):
            year = y
        if year is None:
            continue
        period = period_of(sheet, year, str(ws.cell(r, 2).value or "")) if quarterly else period_of(sheet, year, "")
        if period is None:
            continue
        cell = ws.cell(r, col)
        val = cell_text(cell.value, cell.number_format)
        if val is None:  # 「-」（1994 の前期比なし）・空欄は欠測＝取り込まない
            continue
        flag = str(ws.cell(r, flag_col).value or "").strip() if flag_col else ""
        recs.append(ValueRecord(series_id=s.series_id, period=period, region="JP", value=val,
                                status="※（-0.0）" if flag == "※" else "", vintage=published or day,
                                retrieved_at=day, published_at=published,
                                accessor={"type": "cao_gap_xlsx", "url": BASE + fname, "edition": edition,
                                          "sheet": sheet, "col_header": str(acc.get("col_header", "")),
                                          "cell": f"{cell.coordinate}", "cache": cache_rel(cp)}))
    log.info("%s: %s %s 列%d → 値 %d vintage=%s", s.series_id, fname, sheet, col, len(recs), published)
    return finish(s, recs, dry_run=dry_run, exc=CaoGapSourceError)


def main(argv: Optional[list[str]] = None) -> int:
    return run_cli(TYPES, ingest_series, "内閣府 ＧＤＰギャップ・潜在成長率 xlsx 取込（索引から最新公表回を発見）", argv)


if __name__ == "__main__":
    sys.exit(main())
