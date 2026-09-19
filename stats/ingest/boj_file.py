"""
「行＝期間・列＝項目」型の公表 xlsx からの取込（日銀 gap.xlsx、社人研 将来推計人口 等）。設計は docs/データソース選定.md §1。

- 表の形：行＝期間（period_col 列に '1983.1Q'／'1983.1 : 1983.2Q'／2020）、列＝項目（見出し行のラベルで確定。同名列は col_occurrence で指定）。
- accessor.type は boj_file（日銀）／ipss_xlsx（社人研）。取込ロジックは共通。
- 値は ESRI と同じ規則＝**セルの表示書式の桁で四捨五入した文字列**（gap.xlsx は '0.00_' ＝小数2桁。日銀サイトの表示に一致）。
- accessor: {"type": "boj_file", "url": "...gap.xlsx", "sheet": "data1", "col_header": "需給ギャップ", "header_row": 2, "first_data_row": 6}
- 原本は stats/data/cache/boj/<取得日>/。

実行（リポジトリ root）：  python -m stats.ingest.boj_file --all
"""
from __future__ import annotations

import sys
from pathlib import Path
from typing import Optional

from polyarchy_common.logsetup import configure_quiet_logging, get_logger

from stats.core.periods import convert, matches_freq
from stats.core.registry import Series
from stats.core.values import ValueRecord
from stats.ingest._base import SourceError, _norm, cache_rel, cell_text, fetch, finish, run_cli, today

configure_quiet_logging()
log = get_logger("polyarchy.stats.ingest.boj_file")
TYPES = ("boj_file", "ipss_xlsx")


class BojSourceError(SourceError):
    pass


def ingest_series(s: Series, *, dry_run: bool = False, day: Optional[str] = None) -> int:
    import openpyxl
    day = day or today()
    acc = s.accessor
    if acc.get("type") not in TYPES:
        raise ValueError(f"{s.series_id}: accessor.type={acc.get('type')} は boj_file/ipss_xlsx の対象外")
    cp, published = fetch(acc["url"], day=day, kind="ipss" if acc.get("type") == "ipss_xlsx" else "boj")
    wb = openpyxl.load_workbook(cp, data_only=True)
    sheet = acc.get("sheet") or wb.sheetnames[0]
    if sheet not in wb.sheetnames:
        raise BojSourceError(f"{s.series_id}: シート {sheet!r} が無い（{wb.sheetnames}）")
    ws = wb[sheet]
    hrow = int(acc.get("header_row", 2))
    want = _norm(acc["col_header"])
    cols = [c for c in range(1, ws.max_column + 1) if _norm(ws.cell(hrow, c).value) == want]
    occ = int(acc.get("col_occurrence", 0) or 0)
    if occ:
        if len(cols) < occ:
            raise BojSourceError(f"{s.series_id}: 列見出し {acc['col_header']!r} の出現 {occ} 番目が無い（{len(cols)} 列）")
        col = cols[occ - 1]
    elif len(cols) != 1:
        raise BojSourceError(f"{s.series_id}: 列見出し {acc['col_header']!r} が {len(cols)} 列（1列に確定しない・col_occurrence を指定）")
    else:
        col = cols[0]
    pcol = int(acc.get("period_col", 1))
    recs: list[ValueRecord] = []
    for r in range(int(acc.get("first_data_row", 6)), ws.max_row + 1):
        lab = ws.cell(r, pcol).value
        if lab is None or str(lab).strip() == "":
            continue
        if isinstance(lab, float) and lab.is_integer():
            lab = int(lab)
        period = convert(s.period_converter, str(lab))
        if period is None or (s.freq and not matches_freq(period, s.freq)):
            raise BojSourceError(f"{s.series_id}: 変換できない期間ラベル {lab!r}（period_converter={s.period_converter}）")
        c = ws.cell(r, col)
        text = cell_text(c.value, c.number_format)
        if text is None:
            continue
        recs.append(ValueRecord(series_id=s.series_id, period=period, region="JP", value=text, status="",
                                vintage=published or day, retrieved_at=day, published_at=published,
                                accessor={"type": acc["type"], "url": acc["url"], "sheet": sheet, "col_header": acc["col_header"],
                                          **({"col_occurrence": occ} if occ else {}),
                                          "cell": f"{c.column_letter}{r}", "row_label": str(lab), "cache": cache_rel(cp)}))
    if not recs:
        raise BojSourceError(f"{s.series_id}: 値が0件")
    log.info("%s: %s!%s 列%d → 値 %d vintage=%s", s.series_id, Path(acc["url"]).name, sheet, col, len(recs), published)
    return finish(s, recs, dry_run=dry_run, exc=BojSourceError)


def main(argv: Optional[list[str]] = None) -> int:
    return run_cli(TYPES, ingest_series, "日銀・社人研 公表 xlsx 取込", argv)


if __name__ == "__main__":
    sys.exit(main())
