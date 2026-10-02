"""
総務省「統一的な基準による財務書類に関する情報」の都道府県分 Excel からの取込（第 12 弾 第 4 便・2026-10-02）。

- 要件源＝利用側の財政 PL・BS 分析（地方の土地・社会資本・負債を公会計側から引く。利用側は手作業でダウンロードしていた）。
  **都道府県まで**（市区町村は収録しない＝2026-10-02 判断・登録簿の大きさ B23）。
- 取得元＝https://www.soumu.go.jp/iken/kokaikei/{H28|H29|H30|R01|…}_chihou_zaimusyorui.html の「都道府県」の Excel（main_content/<番号>.xlsx）。
  番号は年版ごとに違う＝下の FILES に年版→番号を持つ（2026-10-02 に H28〜R06 の頁で確認）。索引＝https://www.soumu.go.jp/iken/kokaikei/index.html。
  総務省サイトの利用規約＝政府標準利用規約（第 2.0 版）準拠（◎・出典明示）。
- 表＝1 シートに 4 表が縦に並ぶ（一般会計等、全体、連結 の 貸借対照表／行政コスト計算書／純資産変動計算書／資金収支計算書）。
  列＝都道府県ごとに 3 列（一般会計等・全体・連結）。都道府県名は貸借対照表の 4 行目（各 3 列の先頭）・会計の区分は各表の見出し行。
  **都道府県がそろわない年版がある**（H28 は 39 都道府県）＝列は見出しで決める（位置で決めない）。単位は百万円（整数）。'-' は値ではない。
- 行＝科目の**字下げの道筋**（B〜F 列の階層。例「固定資産/有形固定資産/インフラ資産/土地」）で特定する＝同名の科目（土地・その他）が何度も出るため。
  A 列の行（資産合計・純行政コスト 等の合計行）は道筋の根にしない。1 表の中で道筋が 2 行以上に当たればエラー（黙って選ばない）。
- 年版のファイルには当年度と前年度の 2 シート。**値はその年度を当年度とする年版のシート**から採り、そこに無い都道府県・年度だけ翌年版の前年度シートで埋める
  （accessor.sheet に刻む）。シート名の年（H28_／R1_ 等）が年版と合わなければエラー。
- accessor: {"type": "soumu_tokitsu_xlsx", "section": "BS"|"PL"|"NW"|"CF", "path": "固定資産/有形固定資産/インフラ資産/土地",
             "basis": "一般会計等"|"全体"|"連結", "editions": [2016, …]}
- 原本は stats/data/cache/soumu/<取得日>/tokitsu/。

実行（リポジトリ root）：  python -m stats.ingest.soumu_tokitsu --all
"""
from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import Optional

from polyarchy_common.logsetup import configure_quiet_logging, get_logger

from stats.core.dim_vocab import PREF_CODE
from stats.core.periods import ERA_BASE
from stats.core.registry import Series
from stats.core.values import ValueRecord
from stats.ingest._base import SourceError, _norm, cache_rel, fetch, finish, run_cli, today

configure_quiet_logging()
log = get_logger("polyarchy.stats.ingest.soumu_tokitsu")
BASE = "https://www.soumu.go.jp/main_content/"
INDEX = "https://www.soumu.go.jp/iken/kokaikei/index.html"
# 年版（会計年度）→ 都道府県の Excel の番号（各年版の頁の「都道府県」のリンク）
FILES = {2016: "000608806", 2017: "000677374", 2018: "000734548", 2019: "000796757", 2020: "000836160",
         2021: "000902031", 2022: "000969285", 2023: "001031193", 2024: "001091038"}
SECTIONS = (("貸借対照表", "BS"), ("行政コスト計算書", "PL"), ("純資産変動計算書", "NW"), ("資金収支計算書", "CF"))
BASES = ("一般会計等", "全体", "連結")


class SoumuTokitsuError(SourceError):
    pass


def page_of(year: int) -> str:
    era = f"H{year - 1988}" if year < 2019 else f"R{year - 2018:02d}"
    return f"https://www.soumu.go.jp/iken/kokaikei/{era}_chihou_zaimusyorui.html"


def _sheet_year(title: str) -> Optional[int]:
    m = re.match(r"^([HR])(\d+)_", title or "")
    return (ERA_BASE[m.group(1)] + int(m.group(2))) if m else None


def parse_sheet(ws) -> dict:
    """1 シート → {"cells": {(section, path, basis, pref_code): (値文字列, セル)}}。道筋の重複・列見出しの不整合はエラー。"""
    rows: dict[tuple[str, str], list[int]] = {}
    header: dict[str, int] = {}
    sec, stack = None, {}
    for r in range(1, ws.max_row + 1):
        a = _norm(ws.cell(r, 1).value)
        hit = [k for w, k in SECTIONS if w in a]
        if hit:
            sec, stack = hit[0], {}
            continue
        if a == "指標":  # H28 版だけ末尾に「指標」（住民一人当たり資産額 等）の表が続く＝4 表の外＝ここで終わり
            break
        if not sec:
            continue
        cells = [(c, _norm(ws.cell(r, c).value)) for c in range(1, 7) if _norm(ws.cell(r, c).value)]
        if not cells:
            if sec not in header and any(_norm(ws.cell(r, c).value) == "一般会計等" for c in range(7, 10)):
                header[sec] = r
            continue
        lv, lab = cells[-1]
        if lab == "科目":
            continue
        if lv == 1:  # A 列＝合計行。道筋の根にしない
            stack = {}
            rows.setdefault((sec, lab), []).append(r)
            continue
        stack = {k: v for k, v in stack.items() if k < lv}
        stack[lv] = lab
        rows.setdefault((sec, "/".join(stack[k] for k in sorted(stack))), []).append(r)
    if set(header) != {k for _w, k in SECTIONS}:
        raise SoumuTokitsuError(f"{ws.title}: 会計の区分の見出し行が見つからない表がある（{sorted(header)}）")
    # 列：都道府県名（貸借対照表の 4 行目・3 列ごとの先頭）＋会計の区分（各表の見出し行）。4 表で区分の並びが同じことを確かめる
    colmap: dict[int, tuple[str, str]] = {}
    pref = None
    for c in range(7, ws.max_column + 1):
        name = _norm(ws.cell(4, c).value)
        if name:
            if name not in PREF_CODE:
                raise SoumuTokitsuError(f"{ws.title}: 4 行目 {ws.cell(4, c).coordinate} の {name!r} は都道府県名ではない")
            pref = name
        basis = _norm(ws.cell(header["BS"], c).value)
        if not basis:
            continue
        if basis not in BASES or pref is None:
            raise SoumuTokitsuError(f"{ws.title}: 列 {ws.cell(4, c).column_letter} の見出しが不正（{pref!r}・{basis!r}）")
        for sec, hr in header.items():
            if _norm(ws.cell(hr, c).value) != basis:
                raise SoumuTokitsuError(f"{ws.title}: {sec} の見出し行 {hr} の列 {ws.cell(hr, c).column_letter} が {basis} ではない")
        colmap[c] = (pref, basis)
    for (pf, b) in set(colmap.values()):
        if sum(1 for v in colmap.values() if v == (pf, b)) != 1:
            raise SoumuTokitsuError(f"{ws.title}: {pf} の {b} が複数列")
    cells: dict[tuple[str, str, str, str], tuple[str, str]] = {}
    for (sec, path), rs in rows.items():
        if len(rs) != 1:
            cells[(sec, path, "", "")] = ("DUP", ",".join(map(str, rs)))  # 参照されたらエラー
            continue
        r = rs[0]
        for c, (pf, b) in colmap.items():
            v = ws.cell(r, c).value
            if isinstance(v, bool) or not isinstance(v, (int, float)):
                continue  # '-'・空欄は値ではない
            if isinstance(v, float) and not v.is_integer():  # 百万円の整数でない＝参照されたらエラー（読まない行で止めない）
                cells[(sec, path, b, PREF_CODE[pf])] = ("NONINT", f"{ws.cell(r, c).coordinate}={v}")
                continue
            cells[(sec, path, b, PREF_CODE[pf])] = (str(int(v)), ws.cell(r, c).coordinate)
    return {"cells": cells, "paths": set(rows)}


_BOOK: dict[tuple[int, str], dict] = {}


def _book(year: int, day: str) -> dict:
    """年版の Excel → {fy: parse_sheet の結果＋sheet 名}（当年度・前年度の 2 シート）。"""
    import openpyxl  # 取込時のみ
    key = (year, day)
    if key not in _BOOK:
        cp, published = fetch(BASE + FILES[year] + ".xlsx", day=day, kind="soumu", subdir="tokitsu", name=f"tokitsu_pref_fy{year}.xlsx",
                              skip_if_exists=True)
        wb = openpyxl.load_workbook(cp, data_only=True)
        out = {}
        for i, ws in enumerate(wb.worksheets[:2]):
            fy = _sheet_year(ws.title)
            if fy != (year if i == 0 else year - 1):
                raise SoumuTokitsuError(f"FY{year} 版の {i + 1} 枚目のシート {ws.title!r} の年が {year if i == 0 else year - 1} ではない")
            out[fy] = {**parse_sheet(ws), "sheet": ws.title}
        _BOOK[key] = {"sheets": out, "cache": cache_rel(cp), "published": published}
    return _BOOK[key]


def ingest_series(s: Series, *, dry_run: bool = False, day: Optional[str] = None) -> int:
    day = day or today()
    acc = s.accessor
    if acc.get("type") != "soumu_tokitsu_xlsx":
        raise ValueError(f"{s.series_id}: accessor.type={acc.get('type')} は soumu_tokitsu 取込の対象外")
    sec, path, basis = acc["section"], acc["path"], acc["basis"]
    editions = sorted(int(y) for y in acc["editions"])
    books = {y: _book(y, day) for y in editions}
    for y, b in books.items():
        for fy, sh in b["sheets"].items():
            if (sec, path, "", "") in sh["cells"]:
                raise SoumuTokitsuError(f"{s.series_id}: FY{y} 版 {sh['sheet']} で道筋 {sec}:{path} が複数行")
    recs: list[ValueRecord] = []
    found_path = False
    for fy in range(editions[0] - 1, editions[-1] + 1):
        for code in s.region_codes:
            for ed, which in ((fy, fy), (fy + 1, fy)):  # 当年度の年版 → 翌年版の前年度シート
                b = books.get(ed)
                sh = b["sheets"].get(which) if b else None
                if not sh:
                    continue
                found_path |= (sec, path) in sh["paths"]
                hit = sh["cells"].get((sec, path, basis, code))
                if hit is None:
                    continue
                text, cell = hit
                if text == "NONINT":
                    raise SoumuTokitsuError(f"{s.series_id}: FY{ed} 版 {sh['sheet']} の {cell} が整数でない（百万円の表として読めない）")
                recs.append(ValueRecord(
                    series_id=s.series_id, period=f"FY{fy}", region=code, value=text, status="",
                    vintage=b["published"] or day, retrieved_at=day, published_at=b["published"],
                    accessor={"type": "soumu_tokitsu_xlsx", "file": FILES[ed] + ".xlsx", "edition": f"FY{ed}", "sheet": sh["sheet"],
                              "section": sec, "path": path, "basis": basis, "cell": cell, "cache": b["cache"]}))
                break
    if not found_path:
        raise SoumuTokitsuError(f"{s.series_id}: 道筋 {sec}:{path} がどの年版にも無い")
    log.info("%s: 年版 %s → 値 %d", s.series_id, editions, len(recs))
    return finish(s, recs, dry_run=dry_run, exc=SoumuTokitsuError)


def main(argv: Optional[list[str]] = None) -> int:
    return run_cli(("soumu_tokitsu_xlsx",), ingest_series, "総務省 統一的な基準による財務書類（都道府県）Excel 取込", argv)


if __name__ == "__main__":
    sys.exit(main())
