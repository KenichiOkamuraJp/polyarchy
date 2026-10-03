"""M0 の確認シートを Excel（確認シート.xlsx）で作る。正本は同じフォルダの JSON 2 本。

python deliberations/eval/make_xlsx.py   # JSON を直したら作り直す

確認の結果（確認・コメント欄）は Excel の側に書く＝この xlsx は記入用の作業ファイル（git 外）。
原文へのリンクは目録（data/raw/manifest.jsonl）の source_url に #page=N を付けたもの。
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

D = Path(__file__).resolve().parent
RAW = D.parent / "data" / "raw"
A = json.load(open(D / "anchor_questions.json", encoding="utf-8"))
L = json.load(open(D / "attribution_labels.json", encoding="utf-8"))
URL = {}
if (RAW / "manifest.jsonl").exists():
    for line in open(RAW / "manifest.jsonl", encoding="utf-8"):
        r = json.loads(line)
        if r.get("path"):
            URL[r["path"]] = r.get("source_url")

FONT = "Meiryo"
HEAD_FILL = PatternFill("solid", start_color="1F3864")
INPUT_FILL = PatternFill("solid", start_color="FFF2CC")
TODO_FILL = PatternFill("solid", start_color="F8CBAD")  # 今回の確認が要る行（第 1 回の後に変えた）
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def url(path, page=None):
    u = URL.get(path)
    return f"{u}#page={page}" if u and page else u


def carry(item):
    """記入欄の初期値：第 1 回の後に変えた行は空（要確認）、変えていない行は直近の確認結果を写す。"""
    if item.get("change_round2"):
        return [item["change_round2"], "", ""]
    r = (item.get("review") or [{}])[-1]
    return ["", r.get("result") or "", r.get("comment") or ""]


def sheet(wb, title, headers, widths, rows, links=(), inputs=True):
    """rows＝値のリスト（末尾 3 つ＝今回の変更・確認・コメント）。links＝(列番号, 行ごとの URL)。"""
    ws = wb.create_sheet(title)
    ws.append(headers + (["今回の変更", "確認", "コメント"] if inputs else []))
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
    for r in rows:
        ws.append(r)
    link_cols = dict(links)
    for i, row in enumerate(ws.iter_rows(min_row=2), start=0):
        for c in row:
            c.font = Font(name=FONT, size=9)
            c.alignment = WRAP
            c.border = BORDER
        for col, us in link_cols.items():
            if us[i]:
                cell = row[col]
                cell.hyperlink = us[i]
                cell.font = Font(name=FONT, size=9, color="0563C1", underline="single")
        if inputs:
            todo = bool(row[-3].value)
            for c in row[-3:] if todo else row[-2:]:
                c.fill = TODO_FILL if todo else INPUT_FILL
    for k, w in enumerate(widths + ([30, 8, 30] if inputs else [])):
        ws.column_dimensions[chr(ord("A") + k)].width = w
    if inputs:
        n = len(headers) + 1
        dv = DataValidation(type="list", formula1='"✓,✗,保留"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{chr(ord('A') + n)}2:{chr(ord('A') + n)}{len(rows) + 1}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    return ws


wb = Workbook()
ws = wb.active
ws.title = "説明"
intro = [
    ("M0 確認シート（第 1 便の評価問・帰属ラベルの下書き）", True),
    ("作成 2026-10-03。正本は deliberations/eval/ の anchor_questions.json・attribution_labels.json（本ファイルはそこから make_xlsx.py で作った記入用）。", False),
    ("", False),
    ("第 2 回の確認（第 1 回＝2026-10-03 の記入を反映した版）：オレンジの行だけ見てください（第 1 回の後に変えた行＝「今回の変更」に何を変えたか）。", True),
    ("黄色の行は第 1 回の確認結果をそのまま写してあります（変えたいときだけ書き換える）。外した 5 問は「5 外した問」に理由つきで残しています。", False),
    ("", False),
    ("記入するところ：各シートの右端の「確認」「コメント」の 2 列だけ。", True),
    ("・確認＝プルダウンで ✓（このままでよい）／✗（直す）／保留 を選ぶ。", False),
    ("・コメント＝✗・保留のときに理由や直し方を書く（例：「正解は p5 のほうが適切」「問の聞き方が曖昧」）。", False),
    ("", False),
    ("見ていただきたいこと：問として妥当か・正解の文書と発言者（提出者）が問に合っているか・区分が規則どおりか。", True),
    ("根拠の引用とページは、全件を原文の PDF で機械照合済み（引用は抽出した文字列のまま＝行の途中で切れているのは原文の改行位置）。", False),
    ("「原文」の列のリンクは、その PDF を該当ページで開く（ブラウザによってはページ指定が効かない）。", False),
    ("", False),
    ("区分（開発計画 §4.1・2026-10-03 改定）＝会議での役割：", True),
    ("政務＝総理・大臣・官房長官・国家公安委員会委員長・副大臣・政務官／事務局＝その会議の事務局／府省・会議体＝府省・別の政府の会議体・会議の下の研究会", False),
    ("構成員（政府外）＝民間議員・有識者・座長・構成員である自治体の長／外部（ヒアリング）＝構成員でない出席者／不明＝どれにも当たらない（推定で埋めない）", False),
    ("資料の提出者の探し方（第 1 回の確認の後）：資料名 → 表紙 → 本文（連名・ページごとの提出者表示）。どこにも無ければ会議資料の慣行により事務局（根拠＝既定）。束ねた資料はページ単位。他の会議体が作った文書は作成主体を本文に書かれていれば別に持つ。", False),
    ("正解の範囲（第 1 回の確認の後）：議事録・議事要旨の問は、引用の 1 行ではなく発言の全体（見出し・箇条の頭から次の見出し・箇条の頭まで）。その中のどのチャンクを返しても正解。", False),
    ("", False),
    ("問の型：A 政務・回次の特定／B 名前つき要約・逐語（発言者の特定）／C 匿名の要約（発言者は不明で返るべき）／D 決定文書との往復（初めて出た回と文書。発言者は正解にしない）", False),
]
for text, bold in intro:
    ws.append([text])
    ws.cell(ws.max_row, 1).font = Font(name=FONT, bold=bold, size=12 if ws.max_row == 1 else 10)
ws.column_dimensions["A"].width = 140
ws.cell(4, 1).fill = INPUT_FILL

# ── 1. アンカー問
rows, links = [], []
for q in A["questions"]:
    x = q["expected"]
    ev = "\n".join(f"p{v['page']}「{' / '.join(v['quote'])}」" for v in q["evidence"])
    if q["type"].startswith("D"):
        ans = "\n".join(f"{a['path']}（{a['material_no'] or '番号なし'}）" for a in x["acceptable_paths"])
        ans = f"初出 {x['first_date']} {x['org']} 第{x['session_no']}回\n{ans}"
        who = "（正解にしない）許容：" + "・".join(x.get("acceptable_speakers", []))
        link = url(x["acceptable_paths"][0]["path"], q["evidence"][0]["page"])
    else:
        ans = f"{x['path']} p{x['page']}（{x['material_no'] or x.get('doc_kind', '')}）"
        sp = x.get("answer_span")
        if sp:
            ans += f"\n正解の範囲 p{sp['start_page']}–{sp['end_page']}（{sp['lines']} 行）：{sp['first_line'][:30]} … {sp['last_line'][-20:]}"
        who = f"{x.get('speaker') or x.get('presenter')}（{x.get('speaker_role') or x.get('presenter_type')}）"
        link = url(x["path"], x["page"])
    rows.append([q["id"], q["type"][0], q["question"], ans, who, ev, "原文" if link else "", q.get("note") or ""] + carry(q))
    links.append(link)
sheet(wb, "1 アンカー問", ["ID", "型", "問", "正解（ファイル・ページ）", "発言者／提出者（区分）", "根拠の引用", "原文", "注"],
      [6, 4, 40, 30, 24, 50, 6, 30], rows, [(6, links)])

# ── 2. 資料
rows, links = [], []
for m in L["materials"]:
    pres = m["presenter"]
    if m.get("page_presenters"):
        pres += "\n" + "／".join(f"p{k} {v}" for k, v in m["page_presenters"].items())
    origin = m.get("origin_body") or ""
    if m.get("origin_basis"):
        origin += f"\n（{m['origin_basis']}）"
    rows.append([m["id"], f"{m['org']} 第{m['session_no']}回", m["material_no"], m["material_name_in_list"],
                 m["presenter_type"], pres, m.get("presenter_basis") or "", m["rule"], m.get("cover_evidence") or "",
                 origin, "原文" if url(m["path"], 1) else "", m.get("note") or ""] + carry(m))
    links.append(url(m["path"], 1))
sheet(wb, "2 資料の区分", ["ID", "会議・回", "資料番号", "一覧の資料名", "区分", "提出者", "根拠", "当たった規則", "表紙の根拠の行", "作成主体", "原文", "注"],
      [6, 14, 9, 36, 14, 26, 8, 34, 26, 26, 6, 30], rows, [(10, links)])

# ── 3. 発言（名前あり）
rows, links = [], []
for u in L["utterances_named"]:
    pg = f"p{u['page']}" + (f"–{u['quote_pages'][-1]}" if u.get("quote_pages") else "")
    rows.append([u["id"], f"{u['org']} 第{u['session_no']}回", pg, u["speaker"], u["speaker_role"], u["rule"],
                 " / ".join(u["quote"]), "原文" if url(u["path"], u["page"]) else "", u.get("note") or ""] + carry(u))
    links.append(url(u["path"], u["page"]))
sheet(wb, "3 発言（名前あり）", ["ID", "会議・回", "ページ", "発言者", "区分", "当たった規則", "引用", "原文", "注"],
      [6, 14, 7, 22, 14, 26, 50, 6, 30], rows, [(7, links)])

# ── 4. 発言（匿名）
rows, links = [], []
for u in L["utterances_anonymous"]:
    rows.append([u["id"], f"{u['org']} 第{u['session_no']}回", f"p{u['page']}", u["rule"], u["heading_in_text"],
                 " / ".join(u["quote"]), "原文" if url(u["path"], u["page"]) else "", u.get("note") or ""] + carry(u))
    links.append(url(u["path"], u["page"]))
sheet(wb, "4 発言（匿名）", ["ID", "会議・回", "ページ", "当たった規則", "直前の見出し", "引用", "原文", "注"],
      [6, 14, 7, 30, 36, 50, 6, 30], rows, [(6, links)])

# ── 5. 外した問（記入不要）
rows = [[q["id"], q["question"], q.get("dropped_reason", ""), (q.get("review") or [{}])[-1].get("comment") or ""]
        for q in A.get("dropped", [])]
if rows:
    sheet(wb, "5 外した問", ["ID", "問", "外した理由", "第 1 回のコメント"], [6, 50, 50, 50], rows, inputs=False)

out = D / "確認シート.xlsx"
wb.save(out)
print(out, {s.title: s.max_row - 1 for s in wb.worksheets[1:]})
